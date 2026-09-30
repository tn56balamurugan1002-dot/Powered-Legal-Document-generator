import io,re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF

def sanitize_text(text):
    text = text or ""
    for a,b in {"\u2018":"'","\u2019":"'","\u201c":"\"","\u201d":"\"","\u2013":"-","\u2014":"-","\u2026":"...","\u00a0":" "}.items(): text=text.replace(a,b)
    return text.strip()

def heading(line): return bool(re.match(r"^(ARTICLE|SECTION|CLAUSE)\s+[\w.-]+", line.strip(), re.I) or re.match(r"^\d+[\.)]\s+\S+", line.strip()))

def format_docx(text, doc_type):
    doc=Document(); doc.styles["Normal"].font.name="Times New Roman"; doc.styles["Normal"].font.size=Pt(12)
    t=doc.add_heading(doc_type or "Legal Document",0); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for line in sanitize_text(text).splitlines():
        line=line.strip()
        if not line: continue
        p=doc.add_paragraph(); r=p.add_run(line); r.font.name="Times New Roman"; r.font.size=Pt(13 if heading(line) else 12); r.bold=heading(line)
    f=doc.sections[0].footer.paragraphs[0]; f.text="LegalEase - AI-Assisted Legal Document Draft"; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    b=io.BytesIO(); doc.save(b); return b.getvalue()

class LegalPDF(FPDF):
    def __init__(self, title): super().__init__(); self.title=title
    def header(self): self.set_font("Times","B",12); self.cell(0,8,self.title[:80],new_x="LMARGIN",new_y="NEXT",align="C"); self.ln(4)
    def footer(self): self.set_y(-15); self.set_font("Times",size=9); self.cell(0,10,"LegalEase - AI-Assisted Legal Document Draft",align="C")

def format_pdf(text, doc_type):
    pdf=LegalPDF(doc_type or "Legal Document"); pdf.set_auto_page_break(True,20); pdf.add_page(); pdf.set_font("Times",size=12)
    for raw in sanitize_text(text).splitlines():
        line=raw.strip()
        if not line: pdf.ln(4); continue
        if heading(line):
            pdf.set_font("Times","B",13)
            pdf.multi_cell(0,8,line,new_x="LMARGIN",new_y="NEXT")
            pdf.set_font("Times",size=12)
        else:
            pdf.multi_cell(0,7,line,new_x="LMARGIN",new_y="NEXT")
    out=pdf.output(dest="S"); return out.encode("latin-1",errors="replace") if isinstance(out,str) else bytes(out)
